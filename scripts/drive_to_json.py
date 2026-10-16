#!/usr/bin/env python3
"""Auto-pull: Drive workbook -> data.json (one direction only; the sheet is never written).
Reads the xlsx via a Google service account, converts to the site's data.json shape,
preserves updates/images/imageUrls from the existing data.json, writes only if changed.
Fails safe: any error leaves data.json untouched (site keeps last good data)."""
import json, os, sys, datetime
import openpyxl

FILE_ID = os.environ.get('DRIVE_FILE_ID', '1UuO5FEUyUfGIfDMoEuqFwfb7BAwmEzGV')
SHEETS = ['מוצרים', 'מלאי ישן', 'Sheet1']  # 'מידע' intentionally excluded (as in the manual export)
OUT = 'data.json'
# Owner decision 2026-10-16: publish the sheet as-is, including his name in owner columns.

def conv(v):
    if v is None: return ''
    if isinstance(v, float) and v == int(v): return int(v)
    return v

def convert_sheet(ws):
    rows = [[conv(c) for c in r] for r in ws.iter_rows(values_only=True)]
    while rows and all(c == '' for c in rows[-1]): rows.pop()
    if not rows: return []
    ncol = max((max((i+1 for i, c in enumerate(r) if c != ''), default=0)) for r in rows)
    rows = [list(r[:ncol]) + ['']*(ncol-len(r)) for r in rows]
    return rows

def main():
    key = os.environ.get('DRIVE_SA_KEY')
    if not key:
        print('NOTICE: DRIVE_SA_KEY not set; skipping pull (setup pending).')
        return 0
    from google.oauth2 import service_account
    from googleapiclient.discovery import build
    creds = service_account.Credentials.from_service_account_info(
        json.loads(key), scopes=['https://www.googleapis.com/auth/drive.readonly'])
    drive = build('drive', 'v3', credentials=creds, cache_discovery=False)
    meta = drive.files().get(fileId=FILE_ID, fields='modifiedTime,name').execute()
    data = drive.files().get_media(fileId=FILE_ID).execute()
    with open('/tmp/wb.xlsx', 'wb') as f:
        f.write(data)
    wb = openpyxl.load_workbook('/tmp/wb.xlsx', read_only=True, data_only=True)
    missing = [s for s in SHEETS if s not in wb.sheetnames]
    if missing:
        print('ERROR: sheets missing from workbook:', missing); return 1

    out = {}
    if os.path.exists(OUT):
        out = json.load(open(OUT, encoding='utf-8'))
    tabs = []
    for name in SHEETS:
        tabs.append({'name': name, 'rows': convert_sheet(wb[name])})

    # sanity: don't publish a broken conversion
    if not (100 <= len(tabs[0]['rows']) <= 400): 
        print('ERROR: products row count out of bounds:', len(tabs[0]['rows'])); return 1
    if 'מקט' not in [str(h) for h in tabs[0]['rows'][0]]:
        print('ERROR: products header row unexpected'); return 1

    mt = meta.get('modifiedTime', '')
    try:
        dt = datetime.datetime.fromisoformat(mt.replace('Z', '+00:00')).astimezone(
            datetime.timezone(datetime.timedelta(hours=3)))
        stamp = dt.strftime('%d.%m.%Y %H:%M')
    except Exception:
        stamp = mt
    new = {
        'asOf': 'משיכה אוטומטית. קובץ הדרייב עודכן לאחרונה ב-' + stamp,
        'tabs': tabs,
        'updates': out.get('updates', {}),
        'images': out.get('images', []),
        'imageUrls': out.get('imageUrls', {}),
    }
    if os.path.exists(OUT):
        cur = json.load(open(OUT, encoding='utf-8'))
        if cur.get('tabs') == new['tabs']:
            print('no change in sheet data; data.json left as is')
            return 0
    json.dump(new, open(OUT, 'w', encoding='utf-8'), ensure_ascii=False, separators=(',', ':'))
    print('data.json updated from Drive (modified', stamp + ')')
    return 0

if __name__ == '__main__':
    sys.exit(main())
